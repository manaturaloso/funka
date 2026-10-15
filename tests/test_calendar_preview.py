import importlib.util,tempfile,unittest
from pathlib import Path
from datetime import date,timedelta
import openpyxl
spec=importlib.util.spec_from_file_location('calendar_reader',Path(__file__).parents[1]/'scripts/preview-calendar.py');reader=importlib.util.module_from_spec(spec);spec.loader.exec_module(reader)
class CalendarTests(unittest.TestCase):
 def workbook(self):
  w=openpyxl.Workbook();s=w.active;s.title='Brand';s['A3']=2026;current=date(2026,1,1);starts={}
  for i in range(365):
   d=current+timedelta(days=i);col=i+5;s.cell(5,col,d.day);s.cell(6,col,['L','M','X','J','V','S','D'][d.weekday()]);starts.setdefault(d.month,col)
  for month,col in starts.items():
   end=starts.get(month+1,370)-1;s.cell(3,col,reader.MONTHS[month-1]);s.merge_cells(start_row=3,start_column=col,end_row=3,end_column=end)
  s.cell(8,34,'Promo WP');s.merge_cells(start_row=8,start_column=34,end_row=9,end_column=37);s.cell(11,50,'Branding ambiguo')
  return w
 def test_shifted_headers_merged_span_and_weekdays(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'source.xlsx';self.workbook().save(p);raw=p.read_bytes();data=reader.extract(p);self.assertEqual(data['year'],2026);self.assertEqual(data['sheets'][0]['monthRow'],3);self.assertEqual(data['sheets'][0]['dayRow'],5);self.assertEqual(len(data['campaigns']),2);promo=data['campaigns'][0];self.assertEqual((promo['start'],promo['end']),('2026-01-30','2026-02-02'));self.assertEqual(promo['channelHints'],['WP']);self.assertTrue(data['campaigns'][1]['issues']);self.assertEqual(p.read_bytes(),raw)
 def test_commercial_channels_from_merged_labels_and_seasonal(self):
  with tempfile.TemporaryDirectory() as d:
   w=self.workbook();s=w.active;s['B8']='COMERCIAL WEB PROPIA';s.merge_cells('B8:C10');s['B11']='COMERCIAL SIS';s.merge_cells('B11:C13');s.cell(8,34,'Cyber Umbro: 50%');s.cell(11,50,'COLISEUM DAYS: HASTA 50% OFF');s['B14']='SEASONAL';s.merge_cells('B14:C16');s.cell(15,60,'PRO TRAINING');s.merge_cells(start_row=15,start_column=60,end_row=15,end_column=75)
   p=Path(d)/'source.xlsx';w.save(p);rows=reader.extract(p)['campaigns'];self.assertEqual(rows[0]['channelHints'],['WP']);self.assertEqual(rows[1]['channelHints'],['SIS']);self.assertEqual(rows[2]['channelHints'],[]);self.assertEqual(rows[2]['end'],'2026-03-12')
 def test_conflicting_title_channel_requires_review(self):
  with tempfile.TemporaryDirectory() as d:
   w=self.workbook();s=w.active;s['B8']='COMERCIAL SIS';s.merge_cells('B8:C10');p=Path(d)/'source.xlsx';w.save(p);row=reader.extract(p)['campaigns'][0];self.assertEqual(row['channelHints'],['SIS']);self.assertTrue(any('difiere' in issue for issue in row['issues']))
 def test_ambiguous_year_rejected(self):
  with tempfile.TemporaryDirectory() as d:
   w=self.workbook();w.active['A2']=2025;p=Path(d)/'source.xlsx';w.save(p)
   with self.assertRaises(ValueError):reader.extract(p)
 def test_invalid_day_flagged(self):
  with tempfile.TemporaryDirectory() as d:
   w=self.workbook();w.active.cell(5,starts_february:=36,30);p=Path(d)/'source.xlsx';w.save(p);data=reader.extract(p);self.assertTrue(data['axisIssues'])
if __name__=='__main__':unittest.main()
