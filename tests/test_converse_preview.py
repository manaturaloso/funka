import importlib.util,tempfile,unittest
from pathlib import Path
from datetime import datetime
import openpyxl
spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'scripts/preview-converse.py');reader=importlib.util.module_from_spec(spec);spec.loader.exec_module(reader)
class PreviewTests(unittest.TestCase):
 def test_headers_move_separators_dates_and_mega(self):
  with tempfile.TemporaryDirectory() as d:
   w=openpyxl.Workbook();s=w.active;s.title='Actual';s.append(['Notas']);s.append(['ESTADO','TAREA','FECHA','ENLACE','DEADLINE','ASIGNADO A']);s.append(['Pendiente','PENDIENTES JUEVES']);s.append(['Pendiente','Banner',datetime(2026,10,1),'https://mega.nz/folder/abcd#FULL-secret-fragment',datetime(2026,10,3),'Mar y Cesar']);s.append(['ESTADO','TAREA','FECHA','ENLACE','DEADLINE','ASIGNADO A']);s.append(['Pendiente','Mailing',datetime(2026,10,1),None,'#VALUE!']);s['E6'].data_type='e';f=Path(d)/'source.xlsx';w.save(f);before=f.read_bytes();p=reader.preview(f)
   self.assertEqual(len(p['rows']),2);self.assertEqual(p['rows'][0]['fields']['deadline'],'2026-10-03T00:00:00');self.assertEqual(p['rows'][0]['links'],['https://mega.nz/folder/abcd#FULL-secret-fragment']);self.assertEqual(p['rows'][0]['fields']['responsable'],'Mar y Cesar');self.assertTrue(p['rows'][1]['issues']);self.assertEqual(f.read_bytes(),before)
 def test_deduplication_tracks_conflicting_sources(self):
  with tempfile.TemporaryDirectory() as d:
   w=openpyxl.Workbook();s=w.active;s.title='A';s.append(['FECHA','TAREA','TIPO','ESTADO']);s.append([datetime(2026,10,1),'Banner','WP','Pendiente']);b=w.copy_worksheet(s);b.title='B';c=w.copy_worksheet(s);c.title='C';c['D2']='Entregado';f=Path(d)/'source.xlsx';w.save(f);p=reader.preview(f)
   self.assertEqual(p['summary']['uniqueIdentities'],1);self.assertTrue(p['duplicates'][0]['conflict']);self.assertEqual(len(p['duplicates'][0]['sources']),3);self.assertTrue(all(r['issues'] for r in p['rows']));self.assertEqual(reader.preview(f,['B'])['summary']['candidates'],1)
 def test_blank_historical_date_header_and_unknown_headers(self):
  with tempfile.TemporaryDirectory() as d:
   w=openpyxl.Workbook();s=w.active;s.append([None,'TAREA','TIPO','ESTADO']);s.append([datetime(2025,1,1),'Banner','SIS','Entregado']);w.create_sheet('Notas').append(['No es una tabla']);f=Path(d)/'source.xlsx';w.save(f);p=reader.preview(f);self.assertEqual(p['rows'][0]['fields']['fecha'],'2025-01-01T00:00:00');self.assertFalse(p['sheets'][1]['hasHeaders'])
if __name__=='__main__':unittest.main()
