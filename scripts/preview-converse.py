"""Read-only Converse workbook preview. Never writes to Firebase or the source file."""
import argparse, hashlib, json, re, unicodedata, warnings
from datetime import datetime, date
from pathlib import Path
import openpyxl

ALIASES = {'fecha':['fecha','asignada el','f'], 'tarea':['tarea','descripcion de la tarea'], 'tipo':['tipo'], 'responsable':['asignado a'], 'prioridad':['prioridad'], 'deadline':['deadline'], 'estado':['estado'], 'comentarios':['comentarios'], 'entregado':['entregado el'], 'enlace':['enlace']}
def norm(value):
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFD',str(value or '')) if not unicodedata.combining(c))).strip().lower()
def value(cell):
    v=cell.value
    return v.isoformat() if isinstance(v,(datetime,date)) else v

def preview(path, selected=None):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore',UserWarning)
        book=openpyxl.load_workbook(path,data_only=True)
    rows=[];excluded=[];sheets=[]
    for sheet in book:
        if selected and sheet.title not in selected:continue
        mapping=None;count=0
        for line in sheet.iter_rows():
            headers={field:next((i for i,c in enumerate(line) if norm(c.value) in names),None) for field,names in ALIASES.items()}
            if headers['tarea'] is not None and sum(i is not None for i in headers.values())>=3:
                mapping=headers
                # Some historical sheets omit the date heading but retain date values.
                i=mapping['tarea']
                if mapping['fecha'] is None and i>0 and any(isinstance(sheet.cell(r,i).value,(datetime,date)) for r in range(line[0].row+1,min(sheet.max_row,line[0].row+12)+1)):
                    mapping['fecha']=i-1
                excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'header'});continue
            if not any(c.value is not None for c in line):continue
            if mapping is None:
                excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'before-header'});continue
            title=line[mapping['tarea']]
            if not isinstance(title.value,str) or not title.value.strip() or re.match(r'^pendientes\b',norm(title.value)):
                excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'separator-or-empty-task'});continue
            fields={k:value(line[i]) if i is not None else None for k,i in mapping.items()}
            problems=[]
            for field in ('fecha','deadline','entregado'):
                i=mapping[field]
                if i is not None and line[i].value is not None and not isinstance(line[i].value,(datetime,date)):
                    problems.append(f'{field}: revisar fecha de origen')
            links=[]
            for field in ('enlace','comentarios'):
                i=mapping[field]
                text=str(fields[field] or '')
                links.extend(re.findall(r'https?://[^\s<>"\x00-\x1f]+',text))
                if i is not None and line[i].hyperlink and line[i].hyperlink.target and line[i].hyperlink.target.startswith(('https://','http://')):links.append(line[i].hyperlink.target)
            if norm(fields['estado']) not in ('entregado','aprobado','pendiente','en proceso','en revision','requiere ajustes','cancelado','por asignar',''):
                problems.append('estado: requiere equivalencia')
            identity=[norm(fields['tarea']),norm(fields['tipo']),fields['fecha']]
            # Without a source date, only identical full records are grouped.
            if fields['fecha'] is None:identity.append(fields)
            source_key=hashlib.sha256(json.dumps(identity,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
            rows.append({'id':'converse-'+source_key,'sheet':sheet.title,'row':line[0].row,'brand':'Converse','fields':fields,'links':list(dict.fromkeys(links)),'issues':problems});count+=1
        sheets.append({'name':sheet.title,'candidates':count,'hasHeaders':mapping is not None})
    groups={}
    for row in rows:groups.setdefault(row['id'],[]).append(row)
    duplicates=[]
    for key,items in groups.items():
        if len(items)<2:continue
        versions={json.dumps(r['fields'],sort_keys=True,ensure_ascii=False) for r in items}
        duplicates.append({'id':key,'conflict':len(versions)>1,'sources':[{'sheet':r['sheet'],'row':r['row']} for r in items]})
        if len(versions)>1:
            for r in items:r['issues'].append('identidad repetida con datos distintos: resolver antes de importar')
    return {'format':'funka-converse-preview-v1','sourceSha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),'sheets':sheets,'summary':{'candidates':len(rows),'uniqueIdentities':len(groups),'duplicateGroups':len(duplicates),'rowsWithIssues':sum(bool(r['issues']) for r in rows),'excluded':len(excluded)},'duplicates':duplicates,'excluded':excluded,'rows':rows}
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('file');p.add_argument('--sheet',action='append');p.add_argument('--output',required=True);a=p.parse_args()
    result=preview(a.file,a.sheet)
    if a.sheet and set(a.sheet)!={s['name'] for s in result['sheets']}:p.error('Pestaña inexistente; no se generó salida.')
    out=Path(a.output);out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(result,ensure_ascii=False,indent=2));out.chmod(0o600)
    print(json.dumps(result['summary']))
