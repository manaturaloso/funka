"""Geometric, read-only calendar extraction; ambiguous channels remain for review."""
import argparse,hashlib,json,re,unicodedata
from collections import defaultdict
from datetime import date
from pathlib import Path
import openpyxl
MONTHS=['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']
def norm(v):return ''.join(c for c in unicodedata.normalize('NFD',str(v or '').strip().upper()) if not unicodedata.combining(c))
def integer(v):
 try:
  n=float(v);return int(n) if n.is_integer() else None
 except (TypeError,ValueError):return None

def extract(path):
 w=openpyxl.load_workbook(path,data_only=True);years={integer(c.value) for s in w for row in s.iter_rows(min_row=1,max_row=min(s.max_row,8),max_col=min(s.max_column,3)) for c in row if integer(c.value) and 2000<=integer(c.value)<=2100}
 if len(years)!=1:raise ValueError('No se identifica un año único en el calendario; revisar antes de importar.')
 year=years.pop();records=[];summaries=[];issues=[]
 for s in w:
  header=max(range(1,min(s.max_row,20)+1),key=lambda r:sum(norm(c.value) in MONTHS for c in s[r]))
  months={c.column:MONTHS.index(norm(c.value))+1 for c in s[header] if norm(c.value) in MONTHS}
  if len(set(months.values()))<6:issues.append({'sheet':s.title,'reason':'No se reconoce fila de meses'});continue
  dayrow=max(range(header+1,min(s.max_row,header+8)+1),key=lambda r:sum(integer(c.value) is not None and 1<=integer(c.value)<=31 for c in s[r]))
  dates={};month=None
  for c in s[dayrow]:
   if c.column in months:month=months[c.column]
   n=integer(c.value)
   if month and n and 1<=n<=31:
    try:dates[c.column]=date(year,month,n)
    except ValueError:issues.append({'sheet':s.title,'cell':c.coordinate,'reason':'Día fuera del mes'})
  if len(dates)<300:issues.append({'sheet':s.title,'reason':'Eje temporal incompleto; extracción queda para revisión'})
  merge_at={(m.min_row,m.min_col):m for m in s.merged_cells.ranges};start_count=len(records)
  header_rows=set()
  by_row=defaultdict(list)
  for c in s._cells.values():
   if c.row>dayrow and isinstance(c.value,str) and c.value.strip():by_row[c.row].append(norm(c.value))
  for row,texts in by_row.items():
   weekdays=sum(t in ('L','M','X','J','V','S','D') for t in texts)
   if (weekdays>=20 and weekdays>=len(texts)*0.8) or sum(t in MONTHS for t in texts)>=6:header_rows.add(row)
  # Channel comes from the labelled row block, including vertically merged labels.
  # Seasonal/branding labels never imply a commercial channel.
  channel_blocks=[]
  for label in s._cells.values():
   if label.column>=min(dates,default=1) or not isinstance(label.value,str):continue
   label_text=norm(label.value);channels=[]
   if re.search(r'\b(?:COMERCIAL\s+)?WEB\s+PROPIA\b',label_text) or label_text in ('WP','COMERCIAL WP'):channels=['WP']
   elif re.search(r'\bCOMERCIAL\s+SIS\b',label_text) or label_text=='SIS':channels=['SIS']
   if channels:
    extent=merge_at.get((label.row,label.column));channel_blocks.append((label.row,extent.max_row if extent else label.row,channels))
  # Iterating actual cells avoids phantom formatted rows and handles shifted blocks.
  for cell in sorted(s._cells.values(),key=lambda c:(c.row,c.column)):
   if cell.row<=dayrow or cell.row in header_rows or cell.column not in dates or not isinstance(cell.value,str) or not cell.value.strip():continue
   m=merge_at.get((cell.row,cell.column));endcol=m.max_col if m else cell.column;span=range(cell.column,endcol+1);mapped=[dates[c] for c in span if c in dates]
   warnings=[]
   if len(mapped)!=len(span):warnings.append('El bloque incluye columnas sin fecha; revisar extensión')
   if any(b<=a or (b-a).days!=1 for a,b in zip(mapped,mapped[1:])):warnings.append('Eje temporal discontinuo; revisar fechas')
   text=cell.value.strip();title_channels=set(re.findall(r'\b(WP|SIS)\b',norm(text)))
   block_channels={channel for first,last,values in channel_blocks if first<=cell.row<=last for channel in values}
   channels=sorted(block_channels or title_channels)
   if block_channels and title_channels and title_channels!=block_channels:warnings.append('El canal del título difiere del bloque; revisar antes de generar pedidos')
   if not channels:warnings.append('Canal por confirmar; no generar pedidos arbitrarios')
   source_range=str(m) if m else cell.coordinate
   identity=[s.title,norm(text),mapped[0].isoformat(),mapped[-1].isoformat(),source_range]
   records.append({'id':'cal-'+hashlib.sha256(json.dumps(identity).encode()).hexdigest(),'sheet':s.title,'range':source_range,'title':text,'start':mapped[0].isoformat(),'end':mapped[-1].isoformat(),'channelHints':channels,'issues':warnings,'requiresReview':True})
  summaries.append({'sheet':s.title,'monthRow':header,'dayRow':dayrow,'dateColumns':len(dates),'blocks':len(records)-start_count})
 return {'format':'funka-calendar-preview-v1','sourceSha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),'year':year,'sheets':summaries,'axisIssues':issues,'campaigns':records}
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('file');p.add_argument('--output',required=True);a=p.parse_args();result=extract(a.file);out=Path(a.output);out.write_text(json.dumps(result,ensure_ascii=False,indent=2));out.chmod(0o600);print(json.dumps({'year':result['year'],'sheets':len(result['sheets']),'campaignBlocks':len(result['campaigns']),'axisIssues':len(result['axisIssues'])}))
