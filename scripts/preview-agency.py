"""Read-only multibrand import preview. Workbook contents are data, never commands."""
import argparse,collections,hashlib,json,re,unicodedata,zipfile
from datetime import date,datetime,time
from pathlib import Path
from urllib.parse import urlsplit
import openpyxl

def norm(v):
 return re.sub(r'\s+',' ',''.join(c for c in unicodedata.normalize('NFKD',str(v or '')) if not unicodedata.combining(c) and unicodedata.category(c)!='Cf')).strip().upper()
ALIASES={
 'brand':['MARCA'],'assignee':['ENCARGADO','RESPONSABLE','ENCARGADOS'],
 'assignedOn':['ASIGNADA EL','ASIGNADO EL'],'deadline':['DEADLINE','FECHA DE ENTREGA'],
 'priority':['PRIORIDAD'],'estimatedMinutes':['TIEMPO APROX.','TIEMPO APROX','TIEMPO DE PIEZA'],
 'campaign':['CAMPAÑA'],'type':['TIPO','TIPO DE PEDIDO'],'description':['DETALLE PEDIDO','DETALLE DE PEDIDO'],
 'status':['ESTADO'],'deliveryUrl':['ENLACE DE ENTREGA'],'socialNetwork':['RRSS','RED SOCIAL'],
 'postingDate':['DÍA DE POSTEO','DIA DE PUBLICACION']}
LOOKUP={norm(label):key for key,labels in ALIASES.items() for label in labels}
BRANDS={'NB':'New Balance','NEW BALANCE':'New Balance','COLISEUM':'Coliseum','CAT':'Caterpillar','CATERPILLAR':'Caterpillar','UMBRO':'Umbro','UMBRO COL':'Umbro','FILA':'Fila','HI-TEC':'Hi-Tec','HITEC':'Hi-Tec','MERRELL':'Merrell','STEVE MADDEN':'Steve Madden','CONVERSE':'Converse'}
STATES={norm(s):s for s in ['Por asignar','Pendiente','En proceso','Esperando material','Por validar','En pausa','Entregado','Por subir']}
STATES.update({'EN REVISION':'Por validar'})
PRIORITIES={norm(s):s for s in ['Alta','Media','Baja']}
TYPES={norm(s):s for s in ['POP','Deck','Vitrina','Gráfica de tienda','Mail','Otros','PPLs','Video']}
def serial(v):return v.isoformat() if isinstance(v,(date,datetime,time)) else v
def digest(v):return hashlib.sha256(json.dumps(v,ensure_ascii=False,sort_keys=True,default=serial).encode()).hexdigest()
def blank(v):return v is None or norm(v) in ('','-','—')
def sheet_brand(name):
 n=norm(name)
 for key in sorted(BRANDS,key=len,reverse=True):
  if n==key or n.startswith(key+' '):return BRANDS[key]
 return None

def day(v,field,issues):
 if blank(v):return None
 if isinstance(v,datetime):return v.date().isoformat()
 if isinstance(v,date):return v.isoformat()
 if isinstance(v,str):
  for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y'):
   try:return datetime.strptime(v.strip(),fmt).date().isoformat()
   except ValueError:pass
 issues.append(field+':fecha-no-reconocida');return None

def minutes(v,issues):
 if blank(v):return None
 if isinstance(v,time):return v.hour*60+v.minute+v.second/60
 if isinstance(v,(int,float)) and not isinstance(v,bool) and 0<=v<=100000:return v
 text=norm(v);m=re.fullmatch(r'(\d+):(\d{2})',text)
 if m and int(m[2])<60:return int(m[1])*60+int(m[2])
 m=re.fullmatch(r'(?:(\d+(?:[.,]\d+)?)\s*(?:H|HORA|HORAS))?\s*(?:(\d+(?:[.,]\d+)?)\s*(?:M|MIN|MINUTO|MINUTOS))?',text)
 if m and any(m.groups()):return float((m[1] or '0').replace(',','.'))*60+float((m[2] or '0').replace(',','.'))
 if re.fullmatch(r'\d+(?:[.,]\d+)?',text):return float(text.replace(',','.'))
 issues.append('estimatedMinutes:tiempo-no-reconocido');return None

def preview(path):
 path=Path(path)
 if path.stat().st_size>32*1024*1024:raise ValueError('El Excel supera 32 MiB.')
 with zipfile.ZipFile(path) as archive:
  if sum(info.file_size for info in archive.infolist())>128*1024*1024:raise ValueError('El contenido descomprimido supera 128 MiB.')
 book=openpyxl.load_workbook(path,data_only=True)
 formulas=openpyxl.load_workbook(path,data_only=False)
 if len(book.worksheets)>50:raise ValueError('Demasiadas pestañas para revisar.')
 rows=[];sheets=[];excluded=[];sheet_groups=collections.defaultdict(list);row_groups=collections.defaultdict(list);identities=collections.defaultdict(list)
 for sheet in book:
  if sheet.max_row>10000 or sheet.max_column>80:raise ValueError('La pestaña supera el límite de revisión.')
  sheet_signature=digest([[(serial(c.value),c.hyperlink.target if c.hyperlink else None,str(formulas[sheet.title][c.coordinate].value) if formulas[sheet.title][c.coordinate].data_type=='f' else None) for c in line] for line in sheet.iter_rows()]);sheet_groups[sheet_signature].append(sheet.title)
  mapping=None;count=0;expected=sheet_brand(sheet.title)
  for line in sheet.iter_rows():
   source_formulas={c.coordinate:str(formulas[sheet.title][c.coordinate].value) for c in line if formulas[sheet.title][c.coordinate].data_type=='f'}
   if not any(not blank(c.value) for c in line) and not source_formulas:continue
   headers=collections.defaultdict(list)
   for i,c in enumerate(line):
    if norm(c.value) in LOOKUP:headers[LOOKUP[norm(c.value)]].append(i)
   if 'description' in headers and len(headers)>=3:
    if any(len(v)>1 for v in headers.values()):raise ValueError('Hay encabezados ambiguos en '+sheet.title)
    mapping={k:v[0] for k,v in headers.items()};excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'header'});continue
   if mapping is None:excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'before-header'});continue
   raw={key:serial(line[i].value) for key,i in mapping.items()};description=raw.get('description');campaign=raw.get('campaign')
   if blank(description) and blank(campaign) and not source_formulas and all(blank(v) for k,v in raw.items() if k not in ('brand','socialNetwork')):excluded.append({'sheet':sheet.title,'row':line[0].row,'reason':'empty-template-row'});continue
   issues=['formulas:revisar-valores-cache'] if source_formulas else [];explicit=raw.get('brand');brand=BRANDS.get(norm(explicit)) if not blank(explicit) else expected
   if not blank(explicit) and not brand:issues.append('brand:marca-no-reconocida')
   if expected and brand and expected!=brand:issues.append('brand:conflicto-con-pestana')
   if not brand:issues.append('brand:sin-marca')
   source_value=lambda key:line[mapping[key]].value if key in mapping else None
   assigned=day(source_value('assignedOn'),'assignedOn',issues);deadline=day(source_value('deadline'),'deadline',issues);posting=day(source_value('postingDate'),'postingDate',issues)
   status=STATES.get(norm(raw.get('status')));priority=PRIORITIES.get(norm(raw.get('priority')));kind=TYPES.get(norm(raw.get('type')))
   for key,value in [('status',status),('priority',priority),('type',kind)]:
    if value is None:issues.append(key+':requiere-equivalencia')
   people=[] if blank(raw.get('assignee')) else [v.strip() for v in re.split(r'[,;/&\n]+',str(raw['assignee'])) if v.strip()]
   if people:issues.append('assignee:confirmar-correos-y-accesos')
   links=[];raw_links=[];link_cell=line[mapping['deliveryUrl']] if 'deliveryUrl' in mapping else None
   for candidate in [raw.get('deliveryUrl'),link_cell.hyperlink.target if link_cell is not None and link_cell.hyperlink else None]:
    if blank(candidate):continue
    text=str(candidate).strip();raw_links.append(text)
    try:valid=urlsplit(text).scheme in ('http','https') and bool(urlsplit(text).hostname) and len(re.findall(r'https?://',text,re.I))==1 and not re.search(r'[\s\x00-\x1f]',text)
    except ValueError:valid=False
    if valid:
     if text not in links:links.append(text)
    else:issues.append('deliveryUrl:enlace-no-valido')
   normalized={'brand':brand,'assigneeNames':people,'assignedOn':assigned,'deadlineDate':deadline,'postingDate':posting,'socialNetwork':raw.get('socialNetwork'),'priority':priority,'estimatedMinutes':minutes(source_value('estimatedMinutes'),issues),'campaignName':None if blank(campaign) else str(campaign).strip(),'type':kind,'description':str(description or '').strip(),'status':status,'deliveryLinks':links}
   signature=digest({'raw':raw,'links':raw_links,'formulas':source_formulas});identity=digest([brand,norm(description),norm(campaign),assigned,norm(raw.get('type'))]);record={'previewId':'agency-'+digest([sheet.title,line[0].row,signature]),'sheet':sheet.title,'row':line[0].row,'sheetBrand':expected,'raw':raw,'rawLinks':raw_links,'sourceFormulas':source_formulas,'normalized':normalized,'recordSignature':signature,'identityKey':identity,'issues':issues,'requiresReview':True};rows.append(record);row_groups[signature].append(record);identities[identity].append(record);count+=1
  sheets.append({'name':sheet.title,'expectedBrand':expected,'rows':count,'hasHeaders':mapping is not None,'signature':sheet_signature,'images':len(sheet._images)})
 duplicates=[{'sheets':names,'requiresDecision':True} for names in sheet_groups.values() if len(names)>1]
 repeated=[{'signature':key,'sources':[{'sheet':r['sheet'],'row':r['row']} for r in values]} for key,values in row_groups.items() if len(values)>1]
 conflicts=[{'identityKey':key,'sources':[{'sheet':r['sheet'],'row':r['row']} for r in values]} for key,values in identities.items() if len({r['recordSignature'] for r in values})>1]
 return {'format':'funka-agency-preview-v1','sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest(),'summary':{'rows':len(rows),'distinctRecords':len(row_groups),'duplicateSheetGroups':len(duplicates),'duplicateRecordGroups':len(repeated),'identityConflicts':len(conflicts),'brandConflicts':sum('brand:conflicto-con-pestana' in r['issues'] for r in rows),'responsibleNames':sorted({name for r in rows for name in r['normalized']['assigneeNames']}),'rowsWithIssues':sum(bool(r['issues']) for r in rows)},'sheets':sheets,'duplicateSheets':duplicates,'duplicateRecords':repeated,'identityConflicts':conflicts,'excluded':excluded,'rows':rows,'notes':['Vista previa sin escrituras ni asignaciones a Firebase.','Los identificadores de vista previa no son identificadores de pedidos.','No se aplican prioridades, tipos, responsables ni marcas ambiguas por suposición.','Las imágenes incrustadas se cuentan; su incorporación requiere revisión separada.']}
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('file');p.add_argument('--output',required=True);a=p.parse_args();result=preview(a.file);out=Path(a.output);out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=serial));out.chmod(0o600);print(json.dumps({'rows':result['summary']['rows'],'distinctRecords':result['summary']['distinctRecords'],'duplicateSheetGroups':result['summary']['duplicateSheetGroups'],'brandConflicts':result['summary']['brandConflicts']},ensure_ascii=False))
